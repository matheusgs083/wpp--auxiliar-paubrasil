from __future__ import annotations

import csv
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                Table, TableStyle)


ROOT = Path(r"C:/Users/cadcom.patos")
CSV_PATH = ROOT / "Downloads/consulta-pedidos_21102026_20102026-190209094.csv"
REVENDAS_PATH = ROOT / "Documents/bot_api/data/dRevendas/dRevendas.xlsx"
SETORES_PATH = ROOT / "Documents/bot_api/data/dSetores/dSetores.csv"
OUT_PATH = ROOT / "Documents/bot_api/output/pdf/exemplo_fechamento_gv_setor.pdf"


def norm(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]+", "", text.upper())


def number(value: object) -> float:
    text = str(value or "").strip().replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def money(value: float) -> str:
    return f"R$ {value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def load_mapping() -> tuple[dict[str, str], dict[str, str], dict[tuple[str, str], str]]:
    revendas: dict[str, str] = {}
    delivery_unbs: dict[str, str] = {}
    wb = load_workbook(REVENDAS_PATH, read_only=True, data_only=True)
    rows = wb.active.iter_rows(values_only=True)
    header = next(rows, ())
    columns = {norm(value): index for index, value in enumerate(header)}
    unb_index = columns.get("UNB")
    name_index = columns.get("NOME")
    puxada_index = columns.get("PUXADA")
    if unb_index is None or name_index is None:
        raise ValueError("dRevendas precisa conter as colunas UNB e NOME.")
    for row in rows:
        unb = row[unb_index] if unb_index < len(row) else None
        name = row[name_index] if name_index < len(row) else None
        if unb is not None:
            revendas[str(unb).strip()] = str(name or "").strip()
            if puxada_index is not None and puxada_index < len(row):
                puxada = str(row[puxada_index] or "").strip()
                if puxada:
                    delivery_unbs[puxada] = str(unb).strip()
    setores: dict[tuple[str, str], str] = {}
    with SETORES_PATH.open("r", encoding="utf-8-sig", newline="") as stream:
        for row in csv.DictReader(stream, delimiter=";"):
            filial = str(row.get("Filial") or "").strip()
            setor = str(row.get("Setor") or "").strip()
            gv = str(row.get("Gv") or "").strip()
            if filial and setor and gv:
                setores[(filial, setor)] = gv
    return revendas, delivery_unbs, setores


def resolve_unb(
    raw_code: str,
    operation: str,
    revendas: dict[str, str],
    delivery_unbs: dict[str, str],
) -> str:
    # The order export's delivery-unit code is a depot code (for example 747530),
    # while dRevendas uses the UNB code. Resolve by the named delivery operation.
    raw = str(raw_code).strip()
    if raw in delivery_unbs:
        return delivery_unbs[raw]
    op = norm(operation)
    for unb, name in revendas.items():
        if norm(name) and norm(name) in op:
            return unb
    return raw


def read_orders() -> list[dict[str, object]]:
    revendas, delivery_unbs, setores = load_mapping()
    with CSV_PATH.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream, delimiter=";")
        headers = {norm(h): h for h in (reader.fieldnames or [])}

        def col(row: dict[str, str], name: str) -> str:
            key = headers.get(norm(name))
            return str(row.get(key, "") if key else "").strip()

        grouped: dict[tuple[str, str, str, str], dict[str, object]] = {}
        for raw in reader:
            # All fields are resolved by header name. CSV column letters are never used.
            operation = col(raw, "Desc. unidade entrega") or col(raw, "Desc. unidade venda")
            if not operation:
                continue
            if norm(col(raw, "Desc. operação")) == norm("VENDA PROD CONSIGNAD"):
                continue
            if norm(col(raw, "Situação atend. item")) == norm("FALTA"):
                continue
            raw_unit = col(raw, "Cód. unidade entrega")
            unb = resolve_unb(raw_unit, operation, revendas, delivery_unbs)
            order = col(raw, "Número pedido")
            sector = col(raw, "Cód. setor")
            city = col(raw, "Desc. município")
            if not order or not raw_unit:
                continue
            key = (unb, sector, order, city)
            item = grouped.setdefault(key, {
                "unb": unb,
                "filial": revendas.get(unb, operation),
                "sector": sector or "—",
                "gv": setores.get((unb, sector), "GV não resolvido"),
                "order": order,
                "client": col(raw, "Nome fantasia") or col(raw, "Nome cliente"),
                "city": city,
                "operation": operation,
                "weight": 0.0,
                "value": 0.0,
            })
            item["weight"] = float(item["weight"]) + number(col(raw, "Peso item"))
            item["value"] = float(item["value"]) + number(col(raw, "Valor líquido item") or col(raw, "Valor total"))
        return list(grouped.values())


def build_pdf(rows: list[dict[str, object]]) -> None:
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(OUT_PATH), pagesize=A4,
        rightMargin=9 * mm, leftMargin=9 * mm,
        topMargin=34 * mm, bottomMargin=12 * mm,
        title="Exemplo de fechamento por GV e setor",
    )
    styles = getSampleStyleSheet()
    accent = colors.HexColor("#3B82F6")
    page_bg = colors.HexColor("#0B1220")
    panel_bg = colors.HexColor("#121D2E")
    panel_alt = colors.HexColor("#172438")
    text = colors.HexColor("#E5EEF7")
    muted = colors.HexColor("#A7B6C8")
    title = ParagraphStyle("Title", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=10,
                           leading=11, textColor=text, alignment=TA_LEFT)
    section = ParagraphStyle("Section", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=9.8,
                             leading=10.7, textColor=accent, spaceBefore=4, spaceAfter=2)
    cell = ParagraphStyle("Cell", parent=styles["Normal"], fontName="Courier-Bold", fontSize=6.7, leading=7.5, textColor=text)
    cell_right = ParagraphStyle("CellRight", parent=cell, alignment=2)
    head = ParagraphStyle("Head", parent=cell, fontSize=7.0, leading=7.8, textColor=text)
    note = ParagraphStyle("Note", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=7.2, leading=8.0, textColor=muted)

    totals = defaultdict(lambda: {"orders": 0, "weight": 0.0, "value": 0.0})
    for row in rows:
        key = (str(row["filial"]), str(row["gv"]), str(row["sector"]))
        totals[key]["orders"] += 1
        totals[key]["weight"] += float(row["weight"])
        totals[key]["value"] += float(row["value"])
    summary = [[Paragraph(x, head) for x in ["Filial", "GV", "Setor", "Pedidos", "Peso (kg)", "Valor"]]]
    for (filial, gv, sector), total in sorted(totals.items()):
        summary.append([
            Paragraph(str(filial), cell), Paragraph(str(gv), cell), Paragraph(str(sector), cell),
            Paragraph(str(total["orders"]), cell_right),
            Paragraph(f'{total["weight"]:,.2f}'.replace(",", "X").replace(".", ",").replace("X", "."), cell_right),
            Paragraph(money(total["value"]), cell_right),
        ])
    table = Table(summary, colWidths=[32 * mm, 25 * mm, 18 * mm, 20 * mm, 25 * mm, 30 * mm], repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#203247")), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#33465E")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [panel_bg, panel_alt]), ("TEXTCOLOR", (0, 0), (-1, -1), text), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (3, 1), (-1, -1), "RIGHT"), ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2)]))
    total_orders = sum(int(item["orders"]) for item in totals.values())
    total_weight = sum(float(item["weight"]) for item in totals.values())
    total_value = sum(float(item["value"]) for item in totals.values())
    indicators = [
        [Paragraph("Indicador", head), Paragraph("Valor", head), Paragraph("Indicador", head), Paragraph("Valor", head)],
        [Paragraph("Pedidos agrupados", cell), Paragraph(str(total_orders), cell_right), Paragraph("Clientes fantasia", cell), Paragraph(str(len({str(r['client']) for r in rows})), cell_right)],
        [Paragraph("Peso total", cell), Paragraph(f'{total_weight:,.2f}'.replace(",", "X").replace(".", ",").replace("X", "."), cell_right), Paragraph("Valor total", cell), Paragraph(money(total_value), cell_right)],
        [Paragraph("Filiais encontradas", cell), Paragraph(str(len({str(r['filial']) for r in rows})), cell_right), Paragraph("GVs resolvidos", cell), Paragraph(str(len({str(r['gv']) for r in rows if str(r['gv']) != 'GV não resolvido'})), cell_right)],
    ]
    indicator_table = Table(indicators, colWidths=[48 * mm, 24 * mm, 48 * mm, 24 * mm])
    indicator_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#203247")), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#33465E")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [panel_bg, panel_alt]), ("TEXTCOLOR", (0, 0), (-1, -1), text), ("ALIGN", (1, 1), (-1, -1), "RIGHT"), ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2)]))
    story = [Paragraph("RESUMO DO FECHAMENTO", section), table, Spacer(1, 4), Paragraph("INDICADORES DE SEPARAÇÃO", section), indicator_table, Spacer(1, 4), Paragraph("PEDIDOS POR FILIAL, GV E SETOR", section), Paragraph("A filial é resolvida pela operação e pela dRevendas; o GV é obtido pela chave UNB + setor. Campos ausentes ficam identificados como GV não resolvido.", note), Spacer(1, 3)]
    detail = [[Paragraph(x, head) for x in ["Filial", "GV", "Setor", "Pedido", "Cliente fantasia", "Cidade", "Peso", "Valor"]]]
    for row in sorted(rows, key=lambda item: (str(item["filial"]), str(item["gv"]), str(item["sector"]), str(item["order"])))[:80]:
        weight = f'{float(row["weight"]):,.2f}'.replace(",", "X").replace(".", ",").replace("X", ".")
        value = money(float(row["value"]))
        detail.append([
            Paragraph(str(row["filial"]), cell), Paragraph(str(row["gv"]), cell), Paragraph(str(row["sector"]), cell),
            Paragraph(str(row["order"]), cell), Paragraph(str(row["client"]), cell), Paragraph(str(row["city"]), cell),
            Paragraph(weight, cell_right), Paragraph(value, cell_right),
        ])
    detail_table = Table(detail, colWidths=[25 * mm, 23 * mm, 13 * mm, 17 * mm, 54 * mm, 28 * mm, 17 * mm, 25 * mm], repeatRows=1)
    detail_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#203247")), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#33465E")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [panel_bg, panel_alt]), ("TEXTCOLOR", (0, 0), (-1, -1), text), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (6, 1), (-1, -1), "RIGHT"), ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2)]))
    story += [detail_table]

    def page_header(canvas, doc_obj):
        width, height = A4
        canvas.saveState()
        canvas.setFillColor(page_bg)
        canvas.rect(0, 0, width, height, fill=1, stroke=0)
        canvas.setFillColor(panel_bg)
        canvas.rect(7 * mm, height - 30 * mm, width - 14 * mm, 23 * mm, fill=1, stroke=0)
        canvas.setStrokeColor(colors.HexColor("#33465E"))
        canvas.rect(7 * mm, height - 30 * mm, width - 14 * mm, 23 * mm, fill=0, stroke=1)
        canvas.setFillColor(text)
        canvas.setFont("Helvetica", 7)
        canvas.drawString(9 * mm, height - 9 * mm, "PW02041R-l-Bot API")
        canvas.drawCentredString(width / 2, height - 9 * mm, "Fechamento de Entregas")
        canvas.drawRightString(width - 9 * mm, height - 9 * mm, "EXEMPLO")
        canvas.setFillColor(muted)
        canvas.setFont("Helvetica", 6.5)
        canvas.drawString(9 * mm, height - 13 * mm, "Distribuidora de Bebidas Pau Brasil LTDA")
        canvas.drawRightString(width - 9 * mm, height - 13 * mm, f"Pag. {doc_obj.page}")
        canvas.drawString(9 * mm, height - 17 * mm, "Rotina: fechamento por filial, GV e setor      Fonte: CSV de pedidos")
        canvas.drawRightString(width - 9 * mm, height - 17 * mm, "Campos por nome")
        canvas.drawString(9 * mm, height - 21 * mm, "Filiais: " + ", ".join(sorted({str(r['filial']) for r in rows})))
        canvas.setStrokeColor(accent)
        canvas.line(9 * mm, height - 28 * mm, width - 9 * mm, height - 28 * mm)
        canvas.restoreState()

    doc.build(story, onFirstPage=page_header, onLaterPages=page_header)


if __name__ == "__main__":
    rows = read_orders()
    build_pdf(rows)
    print(f"Gerado: {OUT_PATH} ({len(rows)} pedidos agrupados)")

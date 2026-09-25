
from __future__ import annotations

import csv
import io
import json
import math
import re
import threading
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta
from itertools import chain
from pathlib import Path
from typing import Any, Iterable

from bot_api.services.liga_entrega_motivos import MOTIVOS_DEVOLUCAO
from bot_api.services.liga_entrega_irismark_aux import IRISMARK_AUXILIARY, IRISMARK_AUXILIARY_PLATES, IRISMARK_CODE

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None  # type: ignore[assignment]

METAS = {"devol": 1.4, "saida": "07:30", "saida_pct": 90, "tempo": 110, "tempo_pern": 100, "km": 10, "check_inicio": "2026-07-13"}
PESOS_MOT = {"devol": 35, "saida": 25, "km": 15, "check": 25}
PESOS_AJD = {"devol": 35, "saida": 25, "km": 15, "check": 25}
MIN_ROTAS = 3
TEMPO_PREV_MAX = 840
MAX_AUXILIARY_BYTES = 25 * 1024 * 1024
# Increment when the serialized dashboard shape or enrichment fallback changes;
# otherwise an older persisted payload can hide newly available report fields.
# Increment when the enrichment rules change so a persisted dashboard built
# with an older rule cannot hide newly linked routes or helpers.
CACHE_VERSION = 13

R030805 = "030805_LIGA"
R031120 = "031120_BOT"
R030224S = "030224_RESUMO_LIGA"
R030224M = "030224_MOTORISTA_LIGA"
R030224A = "030224_AJUDANTE_LIGA"
R030237 = "030237"
R03114902 = "03114902_BOT"
R031129 = "031129_LIGA"
RESP = "PONTOMAIS_ESPELHO"
RCHK = "CHECKLIST_FROTA"
ROUTINES = (R030805, R031120, R030224S, R030224M, R030224A, R030237, R03114902, R031129, RESP, RCHK)

# Mesmo status utilizado no painel Liga entregue em 22/09/2026. Quem estiver
# de férias ou desligado continua monitorado, porém não ocupa posição nem prêmio.
CANONICAL_STATUS = {
    "7213": "ferias", "7328": "desligado", "7358": "ferias", "7417": "desligado",
    "7431": "ferias", "7432": "desligado", "7470": "desligado", "7478": "desligado",
}

# Cadastro de motoristas que compõe a Liga. Mantém o denominador operacional
# estável mesmo quando alguém não aparece no CSV do mês.
CANONICAL_MOTORISTAS = {
    "6096": ("JOSIVALDO GOMES DE OLIVEIRA", "PATOS"), "7016": ("RONALDO DINIZ DOS SANTOS", "PATOS"),
    "7302": ("ADRIANO DINIZ PAULO", "PATOS"), "7351": ("LAIRES MENDES DE SOUSA", "PATOS"),
    "7356": ("CARLOS ALEXANDRE LEITE SILVA", "PATOS"), "7358": ("EVERTON OLIVEIRA DE MORAIS", "PATOS"),
    "7362": ("ROBERTO JORGE ALMEIDA DOS SANTOS JUNIOR", "PATOS"), "7375": ("JOSE VANDERLAN DA SILVA GOMES", "PATOS"),
    "7383": ("PEDRO APRIGIO DOS SANTOS FILHO", "PATOS"), "7392": ("CICERO DELFINO DA COSTA", "PATOS"),
    "7394": ("JOFLE LUILLES CARVALHO LEITE", "PATOS"), "7404": ("VINICIUS MENDES GOMES GONCALVES", "PATOS"),
    "7410": ("LEONARDO VIEIRA DA SILVA", "PATOS"), "7430": ("DIOGO DE MEDEIROS LIMA", "PATOS"),
    "7431": ("TULIO BELO DE LIMA", "PATOS"), "7446": ("COSMO JACKSON MONTEIRO SANTANA", "PATOS"),
    "7459": ("JOSE FRANCILEUDO RODRIGUES", "PATOS"), "7478": ("DANRLEI SATORNO SANTOS", "PATOS"),
    "7479": ("WILLYAN DE LIMA SATORNO", "PATOS"), "9078": ("YAN OLIVEIRA PEREIRA", "SUME"),
    "9083": ("JOSE MARCIO CORDEIRO DE SOUZA", "SUME"), "9085": ("VALDECI SOARES DE LIMA", "SUME"),
    "9087": ("JOSE ROBSON RIBEIRO DO NASCIMENTO", "SUME"), "9097": ("JOSE LUCAS DE OLIVEIRA DUARTE", "SUME"),
}

CANONICAL_AJUDANTES = {
    "1771": ("JOSE ALUIZIO PAULO DE SOUZA", "PATOS"), "6046": ("WALDEMIR DE OLIVEIRA", "PATOS"),
    "7077": ("ALEXSANDRO COSTA DE ARAUJO", "PATOS"), "7213": ("ROMARIO DA SILVA FERREIRA", "PATOS"),
    "7218": ("MARCIO NUNES ALVES", "PATOS"), "7227": ("JUCIELSON DE SOUZA COSTA", "PATOS"),
    "7272": ("UELSON NUNES ALVES", "PATOS"), "7282": ("MAURICIO APOLINARIO NOBREGA FILHO", "PATOS"),
    "7316": ("AVANY NOBREGA ALVES", "PATOS"), "7328": ("ANGELO MARCIEL BARBOSA", "PATOS"),
    "7343": ("JOSE ANTONIO DE MARIA NETO", "PATOS"), "7384": ("FABIO JUNHO MENDES FARIAS", "PATOS"),
    "7387": ("FRANCILEUDO MENDES DA SILVA", "PATOS"), "7395": ("ALAN MORAIS DE SOUSA", "PATOS"),
    "7401": ("MICHELL RODRIGO MACEDO DANTAS", "PATOS"), "7417": ("RONDINELY FELIX MARINHO", "PATOS"),
    "7428": ("EXPEDITO ACASSIO DE ARAUJO ROQUE", "PATOS"), "7432": ("DIEGO BRUNO NOBREGA MATIAS", "PATOS"),
    "7442": ("GABRIEL MORAIS BEZERRA", "PATOS"), "7443": ("ANTONIO DE MEDEIROS BATISTA", "PATOS"),
    "7444": ("JOSE MARCELO GALDINO PEREIRA", "PATOS"), "7447": ("MARCELO BARBOSA LUCENA", "PATOS"),
    "7451": ("RUAN VICTOR LEITE BATISTA", "PATOS"), "7454": ("IRLANDIR FERREIRA DE LIRA", "PATOS"),
    "7461": ("LUAN VICTOR DE SOUSA DANTAS", "PATOS"), "7470": ("GABRIEL BORGES CAVALCANTI", "PATOS"),
    "7472": ("NICOLLAS DA SILVA LUCENA", "PATOS"), "7477": ("MATHEUS DA SILVA ALMEIDA", "PATOS"),
    "7480": ("CARLOS ALBERTO NASCIMENTO DE ARAUJO", "PATOS"), "7484": ("GILMAR SOARES FERREIRA", "PATOS"),
    "7485": ("ANTONIO MARCIO DA SILVA  FILHO", "PATOS"), "7486": ("JOSE ANDERSON OLIVEIRA DOS SANTOS", "PATOS"),
    "7487": ("JOSE HENRIQUE DOURADO DA SILVA", "PATOS"), "7489": ("PEDRO LOURENCO MEDEIROS", "PATOS"),
    "7491": ("FRANCISCO SAMUEL SOARES DA SILVA ROCHA", "PATOS"), "9037": ("LUIZ CARLOS FERREIRA DA SILVA", "SUME"),
    "9062": ("GENILSON JOSE DE SOUSA", "SUME"), "9070": ("LEONALDO NUNES DA SILVA", "SUME"),
    "9076": ("IRISMARK CLEMENTE DE LIRA", "SUME"), "9080": ("CICERO EDUARDO GOMES DA SILVA", "SUME"),
    "9088": ("FRANCISCO FLORENCIO DA SILVA JUNIOR", "SUME"), "9093": ("CASSIO CAUE SILVA OLIVEIRA", "SUME"),
    "9094": ("JACKSON DE LIMA SILVA", "SUME"),
}

CANONICAL_ROSTER = set(CANONICAL_MOTORISTAS) | set(CANONICAL_AJUDANTES)

# O painel consulta este endpoint mais de uma vez durante o carregamento. O
# resultado depende dos lotes e expurgos; guardar o último cálculo evita reler
# CSVs grandes quando nada mudou.
_DASHBOARD_CACHE: dict[tuple[str, str], tuple[str, dict[str, Any]]] = {}
_DASHBOARD_CACHE_LOCK = threading.Lock()
# Os auxiliares não mudam a cada consulta do painel. Reaproveitamos o
# resultado enquanto caminho, tamanho e data de alteração permanecerem iguais.
# Isso evita reabrir o XLSX/CSV a cada atualização da tela.
_AUXILIARY_CACHE_LOCK = threading.Lock()
_ESPELHO_CACHE: dict[tuple[str, int, int], dict[str, str]] = {}
_CHECKLIST_CACHE: dict[tuple[str, int, int, tuple[tuple[str, str], ...]], list[dict[str, str]]] = {}


class LigaEntregaDashboardService:
    def __init__(self, *, report_store: Any, expurgo_service: Any, status_service: Any | None = None, dclientes_query_service: Any | None = None, dprodutos_import_service: Any | None = None) -> None:
        self.report_store = report_store
        self.expurgo_service = expurgo_service
        self.status_service = status_service
        self.dclientes_query_service = dclientes_query_service
        self.dprodutos_import_service = dprodutos_import_service

    def build_dashboard(self, *, competencia: str | None = None) -> dict[str, Any]:
        comp = _clean_comp(competencia) if competencia else self._latest_competencia()
        if not comp:
            return _empty("")
        manifests = self._select_manifests(comp)
        expurgos = self._active_expurgos(comp)
        status_overrides = self._status_overrides(comp)
        cache_key = (str(getattr(self.report_store, "root_dir", id(self.report_store))), comp)
        cache_signature = dashboard_signature(manifests, expurgos, status_overrides)
        with _DASHBOARD_CACHE_LOCK:
            cached = _DASHBOARD_CACHE.get(cache_key)
        if cached and cached[0] == cache_signature:
            return cached[1]
        persisted = self._read_persisted_cache(comp, cache_signature)
        if persisted is not None:
            with _DASHBOARD_CACHE_LOCK:
                _DASHBOARD_CACHE[cache_key] = (cache_signature, persisted)
            return persisted
        rotas: dict[str, dict[str, Any]] = {}
        port: dict[str, dict[str, Any]] = {}
        equipes: dict[str, dict[str, Any]] = {}
        cidades: dict[str, str] = {}
        devols: list[dict[str, Any]] = []
        ajud_devol: dict[str, list[str]] = {}
        ent_m: dict[str, set[str]] = defaultdict(set)
        ent_a: dict[str, set[str]] = defaultdict(set)
        entregas_pdvs: set[str] = set()
        entregas_nfs: set[str] = set()
        entregas_qtde_por_produto: dict[str, float] = defaultdict(float)
        entregas_por_filial: dict[str, dict[str, Any]] = defaultdict(
            lambda: {"pdvs": set(), "nfs": set(), "qtde_por_produto": defaultdict(float)}
        )
        resumo_entregas_hl: dict[str, float] = {}
        ponto: dict[str, str] = {}
        checklist: list[dict[str, str]] = []
        colab: dict[str, dict[str, str]] = {
            code: {"cod": code, "nome": name, "filial": filial, "funcao": "MOTORISTA", "status": canonical_status(code, status_overrides)}
            for code, (name, filial) in CANONICAL_MOTORISTAS.items()
        }
        colab.update({
            code: {"cod": code, "nome": name, "filial": filial, "funcao": "AJUDANTE", "status": canonical_status(code, status_overrides)}
            for code, (name, filial) in CANONICAL_AJUDANTES.items()
        })
        warnings: list[str] = []
        def remember(cod: Any, *, nome: str = "", filial: str = "", funcao: str = "") -> str:
            c = norm_code(cod)
            if c == "0":
                return c
            row = colab.setdefault(c, {"cod": c, "nome": f"COD {c}", "filial": filial, "funcao": funcao, "status": canonical_status(c, status_overrides)})
            if nome and str(row.get("nome") or "").startswith("COD "):
                row["nome"] = clean_name(nome)
            if filial and not row.get("filial"):
                row["filial"] = filial
            if funcao and not row.get("funcao"):
                row["funcao"] = funcao
            return c

        for manifest in manifests.get(R030805, []):
            for file in stored_files(manifest):
                filial = filial_from_name(file["filename"])
                try:
                    for row in rows_from(file["path"]):
                        mot = remember(pick(row, "CdMot", "Motorista"), filial=filial, funcao="MOTORISTA")
                        mapa = norm_mapa(pick(row, "Mapa"))
                        if mot == "0" or not mapa:
                            continue
                        aju: list[str] = []
                        for key in ("CdAju1", "Ajudante 1", "CdAju2", "Ajudante 2"):
                            a = remember(pick(row, key), filial=filial, funcao="AJUDANTE")
                            if a != "0" and a not in aju:
                                aju.append(a)
                        km_real = to_float(pick(row, "KmEntr", "Km Entrada")) - to_float(pick(row, "KmSai", "Km Saida"))
                        km_prev = to_float(pick(row, "KmPrev", "Km Previsto"))
                        km_ok = km_prev > 0 and km_real > 0 and km_real < 2000
                        route_data = to_iso(pick(row, "Data"), fallback=manifest_ref(manifest))
                        placa = str(pick(row, "Placa") or "").strip().upper()
                        if (
                            (comp, route_data, filial.upper(), placa) in IRISMARK_AUXILIARY_PLATES
                            or (comp, route_data, filial.upper(), mot) in IRISMARK_AUXILIARY
                        ) and IRISMARK_CODE not in aju:
                            aju.append(IRISMARK_CODE)
                        rotas[mapa] = {
                            "data": route_data, "mapa": mapa, "filial": filial,
                            "mot": mot, "aju": aju, "km_real": round(km_real, 1) if km_ok else None,
                            "km_prev": round(km_prev, 1) if km_ok else None, "tempo_prev": to_min(pick(row, "TempoPrev", "Tempo Prev")),
                            "hs0805": to_time(pick(row, "HrSai", "Hora Saida")), "he0805": to_time(pick(row, "HrEntr", "Hora Entrada")),
                            "entregas": to_int(pick(row, "Entregas")), "cx_carreg": to_float(pick(row, "CxCarreg")),
                            "cx_entreg": to_float(pick(row, "CxEntreg")), "cidade": "", "src": file["filename"],
                        }
                except Exception as exc:  # noqa: BLE001
                    warnings.append(f"03.08.05 {file['filename']}: {exc}")

        for manifest in manifests.get(R031120, []):
            for file in stored_files(manifest):
                filial = filial_from_name(file["filename"])
                try:
                    for row in rows_from(file["path"]):
                        mapa = norm_mapa(pick(row, "Mapa"))
                        fase = str(pick(row, "Fase") or "").lower()
                        if not mapa:
                            continue
                        item = port.setdefault(mapa, {})
                        event = [to_iso(pick(row, "DtOper", "Data"), fallback=manifest_ref(manifest)), to_time(pick(row, "HrOper", "Hora"))]
                        if fase.startswith("entrada"):
                            item["ent"] = event
                        elif fase.startswith("saida") and event[1]:
                            item["sai"] = event
                            item["mot"] = remember(pick(row, "Motorista", "CdMot"), filial=filial, funcao="MOTORISTA")
                except Exception as exc:  # noqa: BLE001
                    warnings.append(f"03.11.20 {file['filename']}: {exc}")

        for manifest in manifests.get(R030224S, []):
            for file in stored_files(manifest):
                filial_code = filial_code_for_liga(filial_from_name(file["filename"]))
                if filial_code not in {"3", "4"}:
                    continue
                try:
                    for row in rows_from(file["path"]):
                        responsabilidade = clean_name(pick(row, "Responsabilidade"))
                        if "TOTAL GERAL FATURADO" in responsabilidade:
                            resumo_entregas_hl[filial_code] = to_float(pick(row, "Volume"))
                            break
                except Exception as exc:  # noqa: BLE001
                    warnings.append(f"03.02.24 resumo {file['filename']}: {exc}")

        for manifest in manifests.get(R030224A, []):
            for file in stored_files(manifest):
                filial = filial_from_name(file["filename"])
                try:
                    for row in rows_from(file["path"]):
                        nota = str(pick(row, "Nota") or "").strip()
                        if not nota:
                            continue
                        data = to_iso(pick(row, "Data"), fallback=manifest_ref(manifest))
                        key = dev_key(nota, pick(row, "Serie", "Série"), data)
                        ajus: list[str] = []
                        for col in ("Ajudante 1", "Ajudante1", "CdAju1", "Ajudante 2", "Ajudante2", "CdAju2"):
                            a = remember(pick(row, col), filial=filial, funcao="AJUDANTE")
                            if a != "0" and a not in ajus:
                                ajus.append(a)
                        ajud_devol[key] = ajus
                except Exception as exc:  # noqa: BLE001
                    warnings.append(f"03.02.24 ajudante {file['filename']}: {exc}")

        for manifest in manifests.get(R030224M, []):
            for file in stored_files(manifest):
                filial = filial_from_name(file["filename"])
                try:
                    for row in rows_from(file["path"]):
                        nota = str(pick(row, "Nota") or "").strip()
                        if not nota:
                            continue
                        data = to_iso(pick(row, "Data"), fallback=manifest_ref(manifest))
                        key = dev_key(nota, pick(row, "Serie", "Série"), data)
                        mot = remember(pick(row, "Motorista", "CdMot"), nome=pick(row, "Nome Motorista"), filial=filial, funcao="MOTORISTA")
                        devols.append({
                            "chave": key, "filial": filial, "cod": mot, "aju": ajud_devol.get(key, []), "data": data,
                            "nota": nota, "serie": str(pick(row, "Serie", "Série") or "").strip(),
                            "cliente_cod": norm_code(pick(row, "Cod. Cliente", "Cod Cliente", "Cliente")),
                            "cliente": str(pick(row, "Nome Cliente", "Cliente Nome") or "").strip(),
                            "cliente_nome_base": "", "area": str(pick(row, "Area") or "").strip(), "rn": "",
                            "setor": str(pick(row, "Setor") or "").strip(), "gv": "",
                            "rn": str(pick(row, "Setor") or "").strip(),
                            "valor": to_float(pick(row, "Valor")), "volume_hl": to_float(pick(row, "Volume", "Hectolitro", "HL")),
                            "data_devolucao": to_iso(pick(row, "Data Devol.", "Data Devolucao", "Dt Devolucao"), fallback=data),
                            "motivo": str(pick(row, "Desc. Motivo", "Motivo") or "").strip(),
                            "resp": str(pick(row, "Cod. Motivo", "Cod Motivo") or "").strip(),
                            "motivo_operacional": "", "motivo_responsabilidade": "",
                            "placa": str(pick(row, "Placa") or "").strip(), "telefone": str(pick(row, "Telefone") or "").strip(),
                            "usuario": str(pick(row, "Usuario") or "").strip(), "hora": str(pick(row, "Hora") or "").strip(),
                            "excluida": False,
                        })
                except Exception as exc:  # noqa: BLE001
                    warnings.append(f"03.02.24 motorista {file['filename']}: {exc}")

        for manifest in manifests.get(R030237, []):
            for file in stored_files(manifest):
                filial = filial_from_name(file["filename"])
                try:
                    for row in rows_from(file["path"]):
                        if str(pick(row, "Status") or "").strip().upper() == "C":
                            continue
                        cliente = norm_code(pick(row, "Cliente", "Cod Cliente", "Cod. Cliente"))
                        filial_code = filial_code_for_liga(filial)
                        filial_entregas = entregas_por_filial[filial_code]
                        if cliente != "0":
                            entregas_pdvs.add(f"{filial_code}|{cliente}")
                            filial_entregas["pdvs"].add(cliente)
                        nota = norm_code(pick(row, "Nota"))
                        if nota != "0":
                            serie = str(pick(row, "Serie", "Série") or "").strip()
                            entregas_nfs.add(f"{filial_code}|{nota}|{serie}")
                            filial_entregas["nfs"].add(f"{nota}|{serie}")
                        produto = norm_code(pick(row, "Produto", "Código Produto", "Cod Produto"))
                        if produto != "0":
                            quantidade = to_float(pick(row, "Qtde", "Quantidade", "Qtd"))
                            entregas_qtde_por_produto[produto] += quantidade
                            filial_entregas["qtde_por_produto"][produto] += quantidade
                        pdv = f"{cliente}|{to_iso(pick(row, 'Dt. Operacao', 'Dt Operacao', 'Data'), fallback=manifest_ref(manifest))}"
                        mot = remember(pick(row, "Motorista", "CdMot"), filial=filial, funcao="MOTORISTA")
                        if mot != "0":
                            ent_m[mot].add(pdv)
                        for col in ("ajudante-1", "ajudante 1", "Ajudante 1", "CdAju1", "ajudante-2", "ajudante 2", "Ajudante 2", "CdAju2"):
                            a = remember(pick(row, col), filial=filial, funcao="AJUDANTE")
                            if a != "0":
                                ent_a[a].add(pdv)
                except Exception as exc:  # noqa: BLE001
                    warnings.append(f"03.02.37 {file['filename']}: {exc}")

        for manifest in manifests.get(R03114902, []):
            for file in stored_files(manifest):
                try:
                    for row in rows_from(file["path"]):
                        mapa = norm_mapa(pick(row, "Mapa"))
                        cidade = str(pick(row, "Cidade", "Municipio", "Município", "Nome Cidade") or "").strip()
                        if mapa and cidade:
                            cidades[mapa] = cidade
                except Exception as exc:  # noqa: BLE001
                    warnings.append(f"03.11.49.02 {file['filename']}: {exc}")

        for manifest in manifests.get(R031129, []):
            for file in stored_files(manifest):
                filial = filial_from_name(file["filename"])
                try:
                    for row in rows_from(file["path"]):
                        mapa = norm_mapa(pick(row, "Mapa"))
                        if not mapa:
                            continue
                        mot = remember(pick(row, "Motorista", "CdMot"), nome=pick(row, "Nome Motorista"), filial=filial, funcao="MOTORISTA")
                        if mot == "0":
                            continue
                        aju: list[str] = []
                        for code_col, name_col in (("Ajudante 1", "Nome Ajudante 1"), ("Ajudante 2", "Nome Ajudante 2"), ("CdAju1", "Nome Ajudante 1"), ("CdAju2", "Nome Ajudante 2")):
                            a = remember(pick(row, code_col), nome=pick(row, name_col), filial=filial, funcao="AJUDANTE")
                            if a != "0" and a not in aju:
                                aju.append(a)
                        equipes[mapa] = {"data": to_iso(pick(row, "Data"), fallback=manifest_ref(manifest)), "filial": filial, "mot": mot, "aju": aju, "sup": str(pick(row, "Nome Superv. Rota", "Supervisor") or "").strip(), "placa": str(pick(row, "Placa") or "").strip()}
                except Exception as exc:  # noqa: BLE001
                    warnings.append(f"03.11.29 {file['filename']}: {exc}")

        for manifest in manifests.get(RESP, []):
            for file in stored_files(manifest):
                path = file["path"]
                if path.stat().st_size > MAX_AUXILIARY_BYTES:
                    warnings.append(f"Espelho de ponto {file['filename']} ignorado: arquivo acima de 25 MB.")
                    continue
                try:
                    cached_ponto = cached_espelho(path)
                    # O cache torna a atualização barata. Na primeira leitura,
                    # o arquivo dentro do limite é processado integralmente;
                    # não descartamos o Espelho só porque os lotes principais
                    # demoraram mais que um limite global.
                    ponto.update(cached_ponto if cached_ponto is not None else cache_espelho(path))
                except Exception as exc:  # noqa: BLE001
                    warnings.append(f"espelho {file['filename']}: {exc}")
        # O checklist é um XLSX pequeno e é a fonte direta da coluna
        # Checklist. Ele não pode ser descartado pelo timeout reservado ao
        # espelho de ponto, pois isso deixa todos os colaboradores em 0%.
        for manifest in manifests.get(RCHK, []):
            for file in stored_files(manifest):
                if file["path"].stat().st_size > MAX_AUXILIARY_BYTES:
                    warnings.append(f"Checklist {file['filename']} ignorado: arquivo acima de 25 MB.")
                    continue
                try:
                    checklist.extend(cache_checklist(file["path"], colab))
                except Exception as exc:  # noqa: BLE001
                    warnings.append(f"checklist {file['filename']}: {exc}")

        for mapa, equipe in equipes.items():
            rota = rotas.setdefault(mapa, {"data": equipe.get("data") or "", "mapa": mapa, "filial": equipe.get("filial") or "", "mot": equipe.get("mot") or "0", "aju": [], "km_real": None, "km_prev": None, "tempo_prev": None, "hs0805": "", "he0805": "", "entregas": 0, "cidade": "", "src": "03.11.29"})
            route_data = str(equipe.get("data") or rota.get("data") or "")
            route_filial = str(equipe.get("filial") or rota.get("filial") or "").upper()
            route_mot = str(equipe.get("mot") or rota.get("mot") or "")
            route_plate = str(equipe.get("placa") or rota.get("placa") or "").strip().upper()
            auxiliary_aju = [IRISMARK_CODE] if (
                (comp, route_data, route_filial, route_plate) in IRISMARK_AUXILIARY_PLATES
                or (comp, route_data, route_filial, norm_code(route_mot)) in IRISMARK_AUXILIARY
            ) else []
            merged_aju = list(dict.fromkeys([*(rota.get("aju") or []), *(equipe.get("aju") or []), *auxiliary_aju]))
            rota.update({"mot": equipe.get("mot") or rota.get("mot"), "aju": merged_aju, "sup": equipe.get("sup") or "", "placa": equipe.get("placa") or ""})
            # A escala é a fonte da data operacional da equipe. O 03.08.05
            # pode registrar a execução no dia seguinte, então a data da
            # equipe precisa prevalecer quando o mapa existe nos dois lotes.
            if equipe.get("data"):
                rota["data"] = equipe["data"]
            if equipe.get("filial") and not rota.get("filial"):
                rota["filial"] = equipe["filial"]

        if self.dclientes_query_service is not None and devols:
            keys = [(str(item.get("filial") or ""), str(item.get("cliente_cod") or "")) for item in devols]
            try:
                client_map = self.dclientes_query_service.lookup_liga_clientes(keys)
            except Exception:
                client_map = {}
            for item in devols:
                key = (filial_code_for_liga(item.get("filial")), norm_code(item.get("cliente_cod")))
                base = client_map.get(key) or {}
                item["cliente_nome_base"] = str(base.get("cliente_nome") or "").strip()
                item["cliente"] = item["cliente_nome_base"] or item.get("cliente") or item.get("cliente_cod") or ""
                item["setor"] = str(base.get("setor") or item.get("setor") or "").strip()
                item["rn"] = str(base.get("rn") or item.get("rn") or item.get("setor") or "").strip()
                item["area"] = str(base.get("area") or item.get("area") or "").strip()
                item["gv"] = str(base.get("gv") or "").strip()

        for item in devols:
            motivo = MOTIVOS_DEVOLUCAO.get(norm_code(item.get("resp")), {})
            item["motivo_operacional"] = str(motivo.get("motivo") or item.get("motivo") or "").strip()
            item["motivo_responsabilidade"] = str(motivo.get("responsabilidade") or "").strip()

        exp_counts = {"devolucao": 0, "tml": 0, "km": 0, "dispersao": 0}
        for item in expurgos:
            if item.get("tipo") in exp_counts:
                exp_counts[str(item["tipo"])] += 1
            item["aplicados"] = 0
        for dev in devols:
            match = match_dev_exp(dev, expurgos)
            if match:
                dev["excluida"] = True
                dev["expurgo_id"] = str(match.get("id") or "")
                match["aplicados"] = int(match.get("aplicados") or 0) + 1
            owner = colab.get(str(dev.get("cod") or ""), {})
            dev["motorista"] = str(owner.get("nome") or dev.get("cod") or "-")
            dev["ajudantes"] = [str(colab.get(code, {}).get("nome") or code) for code in dev.get("aju", [])]

        rotas_list: list[dict[str, Any]] = []
        for mapa, r0 in rotas.items():
            r = dict(r0)
            p = port.get(mapa, {})
            r["hr_sai"] = (p.get("sai") or [None, r.get("hs0805") or ""])[1]
            if mapa in cidades:
                r["cidade"] = cidades[mapa]
            r["pernoite"] = bool(r.get("tempo_prev") and float(r["tempo_prev"]) > TEMPO_PREV_MAX)
            r["tempo_real"] = tempo_real(r, p, ponto)
            r["tempo_pct"] = pct(float(r["tempo_real"]) / float(r["tempo_prev"]) * 100) if r.get("tempo_real") and r.get("tempo_prev") else None
            exp_s = match_route_exp(r, expurgos, {"tml"})
            exp_k = match_route_exp(r, expurgos, {"km", "dispersao"})
            r["expurgo_saida"] = bool(exp_s)
            r["expurgo_km"] = bool(exp_k)
            if exp_s:
                exp_s["aplicados"] = int(exp_s.get("aplicados") or 0) + 1
            if exp_k:
                exp_k["aplicados"] = int(exp_k.get("aplicados") or 0) + 1
            rotas_list.append(r)
        rotas_list.sort(key=lambda x: (str(x.get("data") or ""), to_int(x.get("mapa"))))

        has_farol = bool(manifests.get(RCHK))
        first_week = is_first_week(rotas_list)
        motoristas, ajudantes = build_rankings(
            rotas_list, port, devols, {k: len(v) for k, v in ent_m.items()}, {k: len(v) for k, v in ent_a.items()},
            checklist, colab, has_farol=has_farol, first_week=first_week,
        )
        fatores_hecto = {}
        if self.dprodutos_import_service is not None and entregas_qtde_por_produto:
            try:
                fatores_hecto = self.dprodutos_import_service.lookup_fatores_hecto(set(entregas_qtde_por_produto))
            except Exception:
                fatores_hecto = {}
        entregas_hl_por_filial = {
            filial_code: sum(
                float(quantidade) * float(fatores_hecto.get(codigo) or 0)
                for codigo, quantidade in dados["qtde_por_produto"].items()
            )
            for filial_code, dados in entregas_por_filial.items()
        }
        entregas_hl_por_filial.update(resumo_entregas_hl)
        entregas_hl = sum(entregas_hl_por_filial.values())
        operacao = build_operacao(
            rotas_list, devols, motoristas, ajudantes,
            entregas_hl=entregas_hl, entregas_pdvs=entregas_pdvs, entregas_nfs=entregas_nfs,
        )
        operacao["filiais"] = {}
        for filial_code, dados in entregas_por_filial.items():
            filial_hl = entregas_hl_por_filial.get(filial_code, 0)
            filial_name = "PATOS" if filial_code == "3" else "SUME" if filial_code == "4" else filial_code
            operacao["filiais"][filial_name] = build_operacao(
                [], [item for item in devols if filial_code_for_liga(item.get("filial")) == filial_code], [], [],
                entregas_hl=filial_hl,
                entregas_pdvs={f"{filial_code}|{item}" for item in dados["pdvs"]},
                entregas_nfs={f"{filial_code}|{item}" for item in dados["nfs"]},
            )
        cobertura = build_cobertura(rotas_list)
        devolucoes_auxiliares = sorted(devols, key=lambda x: (str(x.get("data_devolucao") or x.get("data") or ""), str(x.get("cliente") or ""), str(x.get("nota") or "")), reverse=True)
        active_devols = [item for item in devols if not item.get("excluida")]
        warnings = list(dict.fromkeys(warnings))
        result = {"ok": True, "competencia": comp, "generated_at": datetime.now().isoformat(timespec="seconds"), "summary": {"ready": bool(rotas_list or devols), "rotas": len(rotas_list), "motoristas": len(motoristas), "motoristas_ativos": sum(1 for item in motoristas if item.get("status") == "ativo"), "motoristas_elegiveis": sum(1 for item in motoristas if item.get("elegivel")), "ajudantes": len(ajudantes), "devolucoes": len(active_devols), "devolucoes_expurgadas": len([d for d in devols if d.get("excluida")]), "devolucoes_volume_hl": round(sum(float(d.get("volume_hl") or 0) for d in active_devols), 2), "devolucoes_valor": round(sum(float(d.get("valor") or 0) for d in active_devols), 2), "expurgos": sum(exp_counts.values()), "arquivos": sum(int(m.get("file_count") or 0) for v in manifests.values() for m in v), "warnings": len(warnings)}, "metas": METAS, "pesos": {"motorista": PESOS_MOT, "ajudante": PESOS_AJD}, "reports": manifest_summary(manifests), "rankings": {"motoristas": motoristas, "ajudantes": ajudantes}, "rotas": rotas_list, "devolucoes": devolucoes_auxiliares, "devolucoes_auxiliares": devolucoes_auxiliares, "operacao": operacao, "equipe": sorted(colab.values(), key=lambda x: (x.get("funcao") or "", x.get("nome") or "")), "cobertura": cobertura, "expurgos": {"counts": exp_counts, "items": expurgos}, "first_week": first_week, "has_farol": has_farol, "warnings": warnings[:50]}
        with _DASHBOARD_CACHE_LOCK:
            _DASHBOARD_CACHE[cache_key] = (cache_signature, result)
        self._write_persisted_cache(comp, cache_signature, result)
        return result

    def _cache_path(self, comp: str) -> Path:
        root = Path(getattr(self.report_store, "root_dir", Path("exports/liga_entrega_reports"))).parent / "liga_entrega_dashboard_cache"
        safe_comp = re.sub(r"[^0-9-]", "", comp)
        return root / f"{safe_comp}.json"

    def _read_persisted_cache(self, comp: str, signature: str) -> dict[str, Any] | None:
        try:
            payload = json.loads(self._cache_path(comp).read_text(encoding="utf-8"))
            result = payload.get("result") if isinstance(payload, dict) else None
            return result if payload.get("signature") == signature and isinstance(result, dict) else None
        except (OSError, ValueError, json.JSONDecodeError):
            return None

    def _write_persisted_cache(self, comp: str, signature: str, result: dict[str, Any]) -> None:
        try:
            path = self._cache_path(comp)
            path.parent.mkdir(parents=True, exist_ok=True)
            temp = path.with_suffix(".tmp")
            temp.write_text(json.dumps({"signature": signature, "result": result}, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
            temp.replace(path)
        except OSError:
            return

    def _latest_competencia(self) -> str:
        latest = ""
        for routine in ROUTINES:
            try:
                manifest = self.report_store.latest_manifest(routine)
            except Exception:
                continue
            ref = str((manifest or {}).get("reference_date") or "")
            if len(ref) >= 7 and ref[:7] > latest:
                latest = ref[:7]
        return latest

    def _select_manifests(self, comp: str) -> dict[str, list[dict[str, Any]]]:
        out: dict[str, list[dict[str, Any]]] = {R030805: self._list(R030805, comp)}
        for routine in ROUTINES[1:]:
            items = self._list(routine, comp)
            out[routine] = [items[-1]] if items else []
        return out

    def _list(self, routine: str, comp: str) -> list[dict[str, Any]]:
        if hasattr(self.report_store, "list_manifests"):
            return list(self.report_store.list_manifests(routine, competencia=comp))
        manifest = self.report_store.latest_manifest(routine)
        return [manifest] if manifest and str(manifest.get("reference_date") or "").startswith(comp + "-") else []

    def _active_expurgos(self, comp: str) -> list[dict[str, Any]]:
        try:
            result = self.expurgo_service.list_expurgos(competencia=comp, active_only=True)
        except Exception:
            return []
        return [dict(x) for x in result.get("items", []) if isinstance(x, dict)]

    def _status_overrides(self, comp: str) -> dict[str, str]:
        try:
            return dict(self.status_service.statuses(comp)) if self.status_service is not None else {}
        except Exception:
            return {}



def _empty(comp: str) -> dict[str, Any]:
    return {"ok": True, "competencia": comp, "summary": {"ready": False, "rotas": 0, "motoristas": 0, "ajudantes": 0, "devolucoes": 0, "devolucoes_expurgadas": 0, "devolucoes_volume_hl": 0, "devolucoes_valor": 0, "expurgos": 0, "arquivos": 0, "warnings": 0}, "metas": METAS, "pesos": {"motorista": PESOS_MOT, "ajudante": PESOS_AJD}, "reports": {}, "rankings": {"motoristas": [], "ajudantes": []}, "rotas": [], "devolucoes": [], "devolucoes_auxiliares": [], "operacao": {}, "equipe": [], "cobertura": [], "expurgos": {"counts": {"devolucao": 0, "tml": 0, "km": 0, "dispersao": 0}, "items": []}, "warnings": []}


def dashboard_signature(manifests: dict[str, list[dict[str, Any]]], expurgos: list[dict[str, Any]], statuses: dict[str, str] | None = None) -> str:
    """Assinatura barata dos únicos dados que alteram o resultado calculado."""

    batches = tuple(
        (routine, tuple((str(item.get("batch_id") or ""), str(item.get("stored_at") or "")) for item in items))
        for routine, items in sorted(manifests.items())
    )
    exclusions = tuple(
        tuple(sorted((str(key), str(value)) for key, value in item.items()))
        for item in sorted(expurgos, key=lambda item: str(item.get("id") or ""))
    )
    return repr((CACHE_VERSION, batches, exclusions, tuple(sorted((statuses or {}).items()))))


def canonical_status(cod: Any, overrides: dict[str, str] | None = None) -> str:
    code = norm_code(cod)
    if code not in CANONICAL_ROSTER:
        return "desligado"
    status = str((overrides or {}).get(code) or CANONICAL_STATUS.get(code) or "ativo").strip().lower()
    return status if status in {"ativo", "ferias", "afastado", "desligado"} else "ativo"


def filial_code_for_liga(value: Any) -> str:
    text = str(value or "").strip().upper()
    if "PATOS" in text or "0003" in text:
        return "3"
    if "SUME" in text or "SUMÉ" in text or "0004" in text:
        return "4"
    digits = re.sub(r"\D", "", text).lstrip("0")
    return digits or text


def _clean_comp(value: str | None) -> str:
    text = str(value or "").strip()
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", text):
        raise ValueError("Competencia invalida. Use AAAA-MM.")
    return text


def stored_files(manifest: dict[str, Any]) -> Iterable[dict[str, Any]]:
    for item in manifest.get("files", []):
        if isinstance(item, dict):
            path = Path(str(item.get("path") or ""))
            if path.is_file():
                yield {"filename": str(item.get("filename") or path.name), "path": path}


def read_text(path: Path) -> str:
    data = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


def text_encoding(path: Path) -> str:
    """Detecta UTF-8 pelo começo do arquivo; os relatórios legados são cp1252."""

    try:
        with path.open("rb") as handle:
            sample = handle.read(65536)
        sample.decode("utf-8-sig")
        return "utf-8-sig"
    except UnicodeDecodeError:
        return "cp1252"


def rows_from(path: Path) -> Iterable[dict[str, str]]:
    """Lê CSV em fluxo e normaliza os cabeçalhos uma única vez por arquivo."""

    # Cabeçalhos dos relatórios são ASCII. Ler diretamente do arquivo evita
    # duplicar em memória CSVs mensais grandes antes de começar o cálculo.
    with path.open("r", encoding=text_encoding(path), errors="replace", newline="") as handle:
        lines = (line for line in handle if line.strip())
        try:
            head = next(lines)
        except StopIteration:
            return
        delimiter = ";" if head.count(";") >= max(head.count(","), head.count("\t")) else ("\t" if head.count("\t") > head.count(",") else ",")
        reader = csv.DictReader(chain([head], lines), delimiter=delimiter)
        headers = {str(name or ""): norm_header(name) for name in (reader.fieldnames or [])}
        for raw in reader:
            if not any(str(value or "").strip() for value in raw.values()):
                continue
            yield {headers.get(str(key or ""), ""): str(value or "").strip() for key, value in raw.items()}


def norm_header(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]", "", text.lower())


def pick(row: dict[str, str], *names: str) -> str:
    for name in names:
        key = norm_header(name)
        if key in row:
            return row[key]
    for name in names:
        key = norm_header(name)
        for real, value in row.items():
            if key and (key in real or real in key):
                return value
    return ""


def clean_name(value: Any) -> str:
    return " ".join(str(value or "").strip().upper().split())


def short_name(value: str) -> str:
    parts = [x for x in str(value or "").split() if x]
    if not parts:
        return ""
    conn = {"DE", "DA", "DO", "DOS", "DAS", "E"}
    second = next((x for x in parts[1:] if x.upper() not in conn), parts[1] if len(parts) > 1 else "")
    return " ".join(x.capitalize() for x in ([parts[0], second] if second else [parts[0]]))


def norm_code(value: Any) -> str:
    digits = re.sub(r"\D", "", str(value or "")).lstrip("0")
    return digits or "0"


def norm_mapa(value: Any) -> str:
    return re.sub(r"\D", "", str(value or "")).lstrip("0")


def to_float(value: Any) -> float:
    text = str(value or "0").strip()
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text or 0)
    except ValueError:
        return 0.0


def to_int(value: Any) -> int:
    try:
        return int(float(str(value or "0").replace(",", ".")))
    except ValueError:
        return 0


def to_min(value: Any) -> int | None:
    text = str(value or "").strip()
    m = re.match(r"^(\d+):(\d{2})", text)
    if m:
        return int(m.group(1)) * 60 + int(m.group(2))
    return int(text) if text.isdigit() else None


def to_time(value: Any) -> str:
    m = re.match(r"^(\d{1,2}):(\d{2})", str(value or "").strip())
    return f"{int(m.group(1)):02d}:{int(m.group(2)):02d}" if m else ""


def to_iso(value: Any, *, fallback: str | None = None) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    if not text:
        return fallback or ""
    for fmt in ("%d/%m/%Y", "%d%m%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text[:10] if fmt == "%Y-%m-%d" else text, fmt).date().isoformat()
        except ValueError:
            pass
    return fallback or ""


def manifest_ref(manifest: dict[str, Any]) -> str:
    return to_iso(manifest.get("reference_date"))


def filial_from_name(filename: str) -> str:
    upper = filename.upper()
    if "PATOS" in upper or "0003" in upper:
        return "PATOS"
    if "SUME" in upper or "SUMÉ" in upper or "0004" in upper:
        return "SUME"
    return ""


def dev_key(nota: Any, serie: Any, data: Any) -> str:
    return f"{str(nota or '').strip()}|{str(serie or '').strip()}|{to_iso(data)}"


def auxiliary_file_key(path: Path) -> tuple[str, int, int]:
    stat = path.stat()
    return (str(path.resolve()), int(stat.st_size), int(stat.st_mtime_ns))


def cached_espelho(path: Path) -> dict[str, str] | None:
    key = auxiliary_file_key(path)
    with _AUXILIARY_CACHE_LOCK:
        value = _ESPELHO_CACHE.get(key)
    return dict(value) if value is not None else None


def cache_espelho(path: Path) -> dict[str, str]:
    key = auxiliary_file_key(path)
    with _AUXILIARY_CACHE_LOCK:
        value = _ESPELHO_CACHE.get(key)
    if value is None:
        value = parse_espelho(path)
        with _AUXILIARY_CACHE_LOCK:
            _ESPELHO_CACHE[key] = dict(value)
    return dict(value)


def cache_checklist(path: Path, colab: dict[str, dict[str, str]]) -> list[dict[str, str]]:
    roster_key = tuple(sorted(
        (str(code), norm_name(row.get("nome")))
        for code, row in colab.items()
        if row.get("nome") and not str(row.get("nome")).startswith("COD ")
    ))
    key = (*auxiliary_file_key(path), roster_key)
    with _AUXILIARY_CACHE_LOCK:
        value = _CHECKLIST_CACHE.get(key)
    if value is None:
        value = parse_checklist(path, colab)
        with _AUXILIARY_CACHE_LOCK:
            _CHECKLIST_CACHE[key] = [dict(row) for row in value]
    return [dict(row) for row in value]



def parse_espelho(path: Path) -> dict[str, str]:
    out: dict[str, str] = {}
    idx_mat = None
    with path.open("r", encoding=text_encoding(path), errors="replace", newline="") as handle:
        for row in csv.reader(handle):
            if not row:
                continue
            if row[0].strip() == "Data":
                idx_mat = next((i for i, v in enumerate(row) if "matr" in v.lower()), None)
                continue
            m = re.search(r"\d{2}/\d{2}/\d{4}", row[0])
            if idx_mat is None or not m or idx_mat >= len(row):
                continue
            cod = norm_code(row[idx_mat])
            horas = sorted(v.strip() for v in row[2:10] if re.fullmatch(r"\d{2}:\d{2}", v.strip()))
            if cod != "0" and horas:
                key = f"{cod}|{to_iso(m.group(0))}"
                out[key] = max(out.get(key, ""), horas[-1])
    return out


def parse_checklist(path: Path, colab: dict[str, dict[str, str]]) -> list[dict[str, str]]:
    if load_workbook is None:
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = wb.active.iter_rows(values_only=True)
    try:
        first_row = next(rows)
    except StopIteration:
        return []
    header = [str(x or "") for x in first_row]
    idx_dt = header_idx(header, "data conclus")
    idx_ex = header_idx(header, "executor")
    idx_tp = header_idx(header, "tipo")
    if idx_dt is None or idx_ex is None:
        return []
    name_to_cod = {norm_name(v.get("nome")): k for k, v in colab.items() if v.get("nome") and not str(v.get("nome")).startswith("COD ")}
    # O export do Checklist já apresentou este erro de digitação no nome do
    # colaborador. Mantemos o vínculo com o cadastro canônico sem alterar o
    # arquivo original enviado pelo usuário.
    checklist_name_aliases = {
        "JOSE LUCAS DE OLIVEIRA DUARDA": "JOSE LUCAS DE OLIVEIRA DUARTE",
    }
    for alias, canonical in checklist_name_aliases.items():
        if canonical in name_to_cod:
            name_to_cod[alias] = name_to_cod[canonical]
    out: list[dict[str, str]] = []
    for row in rows:
        cod = name_to_cod.get(norm_name(row[idx_ex] if idx_ex < len(row) else ""))
        data = to_iso(row[idx_dt] if idx_dt < len(row) else "")
        if cod and data:
            tipo_raw = str(row[idx_tp] if idx_tp is not None and idx_tp < len(row) else "").lower()
            out.append({"cod": cod, "data": data, "tipo": "R" if tipo_raw.startswith("ret") else "S"})
    return out


def header_idx(header: list[str], needle: str) -> int | None:
    key = norm_header(needle)
    return next((i for i, value in enumerate(header) if key in norm_header(value)), None)


def norm_name(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return " ".join(text.upper().split())


def tempo_real(rota: dict[str, Any], port: dict[str, Any], ponto: dict[str, str]) -> int | None:
    out = None
    if port.get("sai") and port.get("ent"):
        t0, t1 = dt_min(port["sai"][0], port["sai"][1]), dt_min(port["ent"][0], port["ent"][1])
        if t0 is not None and t1 is not None and 60 < t1 - t0 < 7200:
            out = t1 - t0
    if out is not None and out > 4320 and port.get("sai"):
        ult = ponto.get(f"{norm_code(rota.get('mot'))}|{port['sai'][0]}")
        a, b = to_min(ult), to_min(port["sai"][1])
        out = a - b if a is not None and b is not None and a > b else None
    if out is None and not rota.get("pernoite"):
        hs, he = to_min(rota.get("hs0805")), to_min(rota.get("he0805"))
        if hs is not None and he is not None:
            delta = he - hs
            if delta < 0:
                delta += 1440
            if 60 < delta <= 1200:
                out = delta
    return out


def dt_min(data: Any, hora: Any) -> int | None:
    iso, mins = to_iso(data), to_min(hora)
    if not iso or mins is None:
        return None
    d = date.fromisoformat(iso)
    return int(datetime(d.year, d.month, d.day).timestamp() // 60) + mins


def pct(value: float) -> float:
    return 0.0 if math.isnan(value) or math.isinf(value) else round(value, 1)


def faixa(pior: float | None) -> float:
    p = float(pior or 0)
    return 1 if p <= 0 else 0.7 if p <= 25 else 0.4 if p <= 50 else 0


def pior_pct(value: float | None, meta: float, *, menor: bool) -> float:
    if value is None:
        return 0
    if meta == 0:
        return 0 if value <= 0 else 100
    return max(0, ((value - meta) if menor else (meta - value)) / meta * 100)


def blank() -> dict[str, float]:
    return {"rotas": 0, "kmR": 0, "kmP": 0, "tR": 0, "tP": 0, "saiOk": 0, "saiTot": 0, "exp_km": 0, "exp_tml": 0}


def build_rankings(rotas: list[dict[str, Any]], port: dict[str, dict[str, Any]], devols: list[dict[str, Any]], ent_m: dict[str, int], ent_a: dict[str, int], checklist: list[dict[str, str]], colab: dict[str, dict[str, str]], *, has_farol: bool, first_week: bool) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    agg_m: dict[str, dict[str, float]] = defaultdict(blank)
    agg_a: dict[str, dict[str, float]] = defaultdict(blank)
    chk_set = {f"{c.get('cod')}|{c.get('data')}|{c.get('tipo')}" for c in checklist}
    rotas_by_mapa = {str(item.get("mapa") or ""): item for item in rotas}
    chk_e: dict[str, int] = defaultdict(int)
    chk_f: dict[str, int] = defaultdict(int)
    for r in rotas:
        mot = norm_code(r.get("mot"))
        if mot != "0":
            agg_m[mot]["rotas"] += 1
            agg_m[mot]["exp_km"] += int(bool(r.get("expurgo_km")))
            agg_m[mot]["exp_tml"] += int(bool(r.get("expurgo_saida")))
        km_ok = r.get("km_real") is not None and r.get("km_prev") is not None and not r.get("expurgo_km")
        if mot != "0" and km_ok:
            agg_m[mot]["kmR"] += float(r.get("km_real") or 0)
            agg_m[mot]["kmP"] += float(r.get("km_prev") or 0)
        t_ok = r.get("tempo_real") is not None and r.get("tempo_prev")
        target = float(r.get("tempo_prev") or 0) * (float(METAS["tempo_pern"] if r.get("pernoite") else METAS["tempo"]) / 100) if t_ok else 0
        if mot != "0" and t_ok:
            agg_m[mot]["tR"] += float(r.get("tempo_real") or 0)
            agg_m[mot]["tP"] += target
        for a0 in r.get("aju") or []:
            a = norm_code(a0)
            if a == "0":
                continue
            agg_a[a]["rotas"] += 1
            agg_a[a]["exp_km"] += int(bool(r.get("expurgo_km")))
            agg_a[a]["exp_tml"] += int(bool(r.get("expurgo_saida")))
            if km_ok:
                agg_a[a]["kmR"] += float(r.get("km_real") or 0)
                agg_a[a]["kmP"] += float(r.get("km_prev") or 0)
            if t_ok:
                agg_a[a]["tR"] += float(r.get("tempo_real") or 0)
                agg_a[a]["tP"] += target
            p = port.get(str(r.get("mapa") or ""))
            if p and p.get("sai") and not r.get("expurgo_saida"):
                agg_a[a]["saiTot"] += 1
                agg_a[a]["saiOk"] += 1 if str(p["sai"][1]) <= str(METAS["saida"]) else 0
        if has_farol and str(r.get("data") or "") >= str(METAS["check_inicio"]) and mot != "0":
            p = port.get(str(r.get("mapa") or ""), {})
            ds = (p.get("sai") or [r.get("data")])[0]
            de = (p.get("ent") or [r.get("data")])[0]
            fez_s = f"{mot}|{ds}|S" in chk_set
            fez_r = f"{mot}|{de}|R" in chk_set or f"{mot}|{next_day(de)}|R" in chk_set
            for c in [mot] + [norm_code(x) for x in (r.get("aju") or []) if norm_code(x) != "0"]:
                chk_e[c] += 2
                chk_f[c] += int(fez_s) + int(fez_r)
    for mapa, p in port.items():
        mot = norm_code(p.get("mot"))
        if mot == "0" or not p.get("sai"):
            continue
        rota = rotas_by_mapa.get(str(mapa), {})
        if rota.get("expurgo_saida"):
            continue
        agg_m[mot]["saiTot"] += 1
        agg_m[mot]["saiOk"] += 1 if str(p["sai"][1]) <= str(METAS["saida"]) else 0
    devol_m, devol_a = aggregate_devolucoes(devols)
    exp_dev_m, exp_dev_a = aggregate_devolucoes(devols, only_expurgadas=True)
    return (
        mount(colab, agg_m, ent_m, devol_m, exp_dev_m, chk_e, chk_f, "MOTORISTA", has_farol=has_farol, first_week=first_week),
        mount(colab, agg_a, ent_a, devol_a, exp_dev_a, chk_e, chk_f, "AJUDANTE", has_farol=has_farol, first_week=first_week),
    )



def aggregate_devolucoes(devols: list[dict[str, Any]], *, only_expurgadas: bool = False) -> tuple[dict[str, int], dict[str, int]]:
    """Conta devoluções por pessoa uma vez, sem varrer a lista por colaborador."""

    motoristas: dict[str, set[str]] = defaultdict(set)
    ajudantes: dict[str, set[str]] = defaultdict(set)
    for item in devols:
        if bool(item.get("excluida")) != only_expurgadas:
            continue
        pair = f"{item.get('cliente_cod')}|{item.get('data')}"
        mot = norm_code(item.get("cod"))
        if mot != "0":
            motoristas[mot].add(pair)
        for helper in item.get("aju") or []:
            cod = norm_code(helper)
            if cod != "0":
                ajudantes[cod].add(pair)
    return ({code: len(pairs) for code, pairs in motoristas.items()}, {code: len(pairs) for code, pairs in ajudantes.items()})


def mount(colab: dict[str, dict[str, str]], agg: dict[str, dict[str, float]], entregas: dict[str, int], devols: dict[str, int], devols_expurgadas: dict[str, int], chk_e: dict[str, int], chk_f: dict[str, int], role: str, *, has_farol: bool, first_week: bool) -> list[dict[str, Any]]:
    pesos = PESOS_MOT if role == "MOTORISTA" else PESOS_AJD
    codes = {k for k, v in colab.items() if v.get("funcao") == role} | set(agg) | set(entregas)
    rows: list[dict[str, Any]] = []
    for cod in sorted(codes):
        info = colab.get(cod, {"nome": f"COD {cod}", "filial": ""})
        g = agg.get(cod, blank())
        ent = int(entregas.get(cod, 0))
        dev = int(devols.get(cod, 0))
        pdev = round(dev / ent * 100, 2) if ent else None
        psaida = pct(g["saiOk"] / g["saiTot"] * 100) if g["saiTot"] else None
        tempo = pct(g["tR"] / g["tP"] * 100) if g["tP"] else None
        # KM abaixo do previsto não é desvio negativo: somente o excedente
        # realizado acima do previsto prejudica a pontuação.
        km = pct(max(0, g["kmR"] - g["kmP"]) / g["kmP"] * 100) if g["kmP"] else None
        # O HTML original concede o peso inteiro enquanto o Farol ainda não
        # foi disponibilizado. Quando existe, mede saída e retorno normalmente.
        check = pct(chk_f[cod] / chk_e[cod] * 100) if chk_e.get(cod) else (100.0 if not has_farol else None)
        pts = {
            "saida": round(pesos["saida"] * faixa(pior_pct(psaida, float(METAS["saida_pct"]), menor=False)), 1) if psaida is not None else 0,
            "devol": round(pesos["devol"] * faixa(max(0, (pdev - float(METAS["devol"])) / float(METAS["devol"]) * 100)), 1) if pdev is not None else 0,
            "km": round(pesos["km"] * faixa(pior_pct(km, float(METAS["km"]), menor=True)), 1) if km is not None else 0,
            "check": round(pesos["check"] * faixa(pior_pct(check, 100, menor=False)), 1) if check is not None else 0,
        }
        measured = {"saida": psaida is not None, "devol": pdev is not None, "km": km is not None, "check": check is not None}
        sw = sum(w for k, w in pesos.items() if measured[k])
        sp = sum(float(pts[k]) for k in pesos if measured[k])
        status = str(info.get("status") or canonical_status(cod))
        if status != "ativo":
            continue
        rows.append({"cod": cod, "nome": info.get("nome") or f"COD {cod}", "nome_zap": short_name(info.get("nome") or f"COD {cod}"), "filial": info.get("filial") or "", "rotas": int(g["rotas"]), "entregas": ent, "devol": dev, "pdev": pdev, "psaida": psaida, "tempo_pct": tempo, "km_desv": km, "check_pct": check, "check_f": chk_f.get(cod) if chk_e.get(cod) else None, "check_e": chk_e.get(cod) or None, "pts": pts, "expurgos": {"devolucao": int(devols_expurgadas.get(cod, 0)), "km": int(g["exp_km"]), "tml": int(g["exp_tml"])}, "total": round(sp / sw * 100, 1) if sw else 0, "status": status, "elegivel": status == "ativo" and (first_week or int(g["rotas"]) >= MIN_ROTAS), "pos": None})
    elig = [x for x in rows if x["elegivel"]]
    elig.sort(key=lambda x: (-float(x.get("total") or 0), x.get("pdev") if x.get("pdev") is not None else 999, -int(x.get("rotas") or 0)))
    for i, row in enumerate(elig, 1):
        row["pos"] = i
    rows.sort(key=lambda x: (x.get("pos") is None, x.get("pos") or 9999, -float(x.get("total") or 0), x.get("nome") or ""))
    return rows


def count_devols(devols: list[dict[str, Any]], cod: str, role: str) -> int:
    pairs = set()
    for d in devols:
        if d.get("excluida"):
            continue
        ok = norm_code(d.get("cod")) == cod if role == "MOTORISTA" else cod in {norm_code(x) for x in d.get("aju") or []}
        if ok:
            pairs.add(f"{d.get('cliente_cod')}|{d.get('data')}")
    return len(pairs)


def match_dev_exp(dev: dict[str, Any], expurgos: list[dict[str, Any]]) -> dict[str, Any] | None:
    for e in expurgos:
        if e.get("tipo") != "devolucao":
            continue
        if e.get("data") and e.get("data") != dev.get("data"):
            continue
        if e.get("filial") and str(e.get("filial")).upper() != str(dev.get("filial")).upper():
            continue
        if str(e.get("escopo") or "individual").lower() == "equipe":
            return e
        cliente = norm_code(e.get("cliente"))
        cod_cliente = norm_code(dev.get("cliente_cod"))
        if cliente != "0" and cliente != cod_cliente:
            continue
        return e
    return None


def match_route_exp(rota: dict[str, Any], expurgos: list[dict[str, Any]], tipos: set[str]) -> dict[str, Any] | None:
    for e in expurgos:
        if e.get("tipo") not in tipos:
            continue
        if str(e.get("escopo") or "individual").lower() != "equipe" and e.get("mapa") and norm_mapa(e.get("mapa")) != norm_mapa(rota.get("mapa")):
            continue
        if e.get("data") and e.get("data") != rota.get("data"):
            continue
        if e.get("filial") and str(e.get("filial")).upper() != str(rota.get("filial")).upper():
            continue
        return e
    return None


def next_day(value: Any) -> str:
    iso = to_iso(value)
    return (date.fromisoformat(iso) + timedelta(days=1)).isoformat() if iso else ""


def build_operacao(
    rotas: list[dict[str, Any]],
    devols: list[dict[str, Any]],
    motoristas: list[dict[str, Any]],
    ajudantes: list[dict[str, Any]],
    *,
    entregas_hl: float = 0,
    entregas_pdvs: set[str] | None = None,
    entregas_nfs: set[str] | None = None,
) -> dict[str, Any]:
    ent = sum(int(r.get("entregas") or 0) for r in rotas)
    dev = len([d for d in devols if not d.get("excluida")])
    saidas = [r for r in rotas if r.get("hr_sai") and not r.get("expurgo_saida")]
    saidas_ok = [r for r in saidas if str(r.get("hr_sai") or "") <= str(METAS["saida"])]
    kms = [r for r in rotas if r.get("km_real") is not None and r.get("km_prev") is not None and not r.get("expurgo_km")]
    kr = sum(float(r.get("km_real") or 0) for r in kms)
    kp = sum(float(r.get("km_prev") or 0) for r in kms)
    devolucoes_volume_hl = sum(float(item.get("volume_hl") or 0) for item in devols)
    devolucoes_pdvs = {
        f"{filial_code_for_liga(item.get('filial'))}|{norm_code(item.get('cliente_cod'))}"
        for item in devols if norm_code(item.get("cliente_cod")) != "0"
    }
    devolucoes_nfs = {
        f"{filial_code_for_liga(item.get('filial'))}|{norm_code(item.get('nota'))}|{str(item.get('serie') or '').strip()}"
        for item in devols if norm_code(item.get("nota")) != "0"
    }
    entregas_pdvs = entregas_pdvs or set()
    entregas_nfs = entregas_nfs or set()
    hl_pct = round(devolucoes_volume_hl / entregas_hl * 100, 2) if entregas_hl > 0 else None
    pdv_pct = round(len(devolucoes_pdvs) / len(entregas_pdvs) * 100, 2) if entregas_pdvs else None
    nf_pct = round(len(devolucoes_nfs) / len(entregas_nfs) * 100, 2) if entregas_nfs else None
    raw_hl_pct = devolucoes_volume_hl / entregas_hl * 100 if entregas_hl > 0 else None
    raw_pdv_pct = len(devolucoes_pdvs) / len(entregas_pdvs) * 100 if entregas_pdvs else None
    total_pct = round(raw_hl_pct + raw_pdv_pct, 2) if raw_hl_pct is not None and raw_pdv_pct is not None else None
    return {
        "entregas": ent, "devolucoes": dev, "devolucao_pct": round(dev / ent * 100, 2) if ent else None,
        "devolucoes_volume_hl": round(devolucoes_volume_hl, 2), "devolucoes_pdvs": len(devolucoes_pdvs),
        "devolucoes_nfs": len(devolucoes_nfs), "entregas_hl": round(entregas_hl, 2),
        "entregas_pdvs": len(entregas_pdvs), "entregas_nfs": len(entregas_nfs),
        "devolucao_hl_pct": hl_pct, "devolucao_pdv_pct": pdv_pct, "devolucao_nf_pct": nf_pct,
        "devolucao_total_pct": total_pct, "devolucoes_valor": round(sum(float(item.get("valor") or 0) for item in devols), 2),
        "saida_pct": pct(len(saidas_ok) / len(saidas) * 100) if saidas else None,
        "km_desv": pct(max(0, kr - kp) / kp * 100) if kp else None, "rotas": len(rotas),
        "motoristas_elegiveis": len([x for x in motoristas if x.get("elegivel")]),
        "ajudantes_elegiveis": len([x for x in ajudantes if x.get("elegivel")]),
    }


def build_cobertura(rotas: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by: dict[str, dict[str, Any]] = {}
    for r in rotas:
        data = str(r.get("data") or "")
        if not data:
            continue
        row = by.setdefault(data, {"data": data, "rotas": 0, "entregas_0805": 0, "mapas": [], "motoristas": set(), "ajudantes": set()})
        row["rotas"] += 1
        row["entregas_0805"] += int(r.get("entregas") or 0)
        row["mapas"].append(str(r.get("mapa") or ""))
        if r.get("mot"):
            row["motoristas"].add(str(r["mot"]))
        row["ajudantes"].update(str(code) for code in r.get("aju") or [] if code)
    return [{**by[k], "motoristas": len(by[k]["motoristas"]), "ajudantes": len(by[k]["ajudantes"])} for k in sorted(by)]


def is_first_week(rotas: list[dict[str, Any]]) -> bool:
    """Até o dia 7, a Liga não aplica o corte mínimo de três rotas."""

    dates = [
        date.fromisoformat(str(item.get("data")))
        for item in rotas
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(item.get("data") or ""))
    ]
    return bool(dates) and max(dates).day <= 7


def manifest_summary(manifests: dict[str, list[dict[str, Any]]]) -> dict[str, list[dict[str, Any]]]:
    out: dict[str, list[dict[str, Any]]] = {}
    for routine, items in manifests.items():
        out[routine] = [{"batch_id": m.get("batch_id"), "reference_date": m.get("reference_date"), "stored_at": m.get("stored_at"), "file_count": m.get("file_count"), "filenames": [f.get("filename") for f in m.get("files", []) if isinstance(f, dict)]} for m in items]
    return out
